import io
import re

def parse_resume(uploaded_file):
    file_type = uploaded_file.name.split(".")[-1].lower()
    try:
        if file_type == "pdf":
            return _parse_pdf(uploaded_file)
        elif file_type == "docx":
            return _parse_docx(uploaded_file)
        elif file_type == "txt":
            return _parse_txt(uploaded_file)
        else:
            return None
    except Exception as e:
        print(f"Parse error: {e}")
        return None

def _parse_pdf(file):
    try:
        import PyPDF2
        reader = PyPDF2.PdfReader(io.BytesIO(file.read()))
        text = ""
        for page in reader.pages:
            extracted = page.extract_text()
            if extracted:
                text += extracted + "\n"
        return _clean_text(text)
    except ImportError:
        return "Please install PyPDF2: pip install PyPDF2"

def _parse_docx(file):
    try:
        from docx import Document
        doc = Document(io.BytesIO(file.read()))
        paragraphs = [para.text for para in doc.paragraphs if para.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text.strip():
                        paragraphs.append(cell.text)
        return _clean_text("\n".join(paragraphs))
    except ImportError:
        return "Please install python-docx: pip install python-docx"

def _parse_txt(file):
    try:
        content = file.read()
        if isinstance(content, bytes):
            content = content.decode("utf-8", errors="ignore")
        return _clean_text(content)
    except Exception as e:
        return f"Error reading file: {e}"

def _clean_text(text):
    if not text:
        return ""
    text = re.sub(r"[^\x20-\x7E\n\t]", " ", text)
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    lines = [line.strip() for line in text.split("\n")]
    text = "\n".join(line for line in lines if line)
    return text.strip()